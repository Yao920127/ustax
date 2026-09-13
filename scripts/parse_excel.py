import openpyxl
import json
import re
import math
from pathlib import Path

# Setup paths
PROJECT_DIR = Path('/Users/yaolo/Desktop/ustax')
EXCEL_FILE = PROJECT_DIR / '2026-State-Individual-Income-Tax-Rates-Brackets.xlsx'
DATA_DIR = PROJECT_DIR / 'data'

STATE_MAP = {
    'Ala.': 'Alabama', 'Alaska': 'Alaska', 'Ariz.': 'Arizona', 'Ark.': 'Arkansas', 'Calif.': 'California', 
    'Colo.': 'Colorado', 'Conn.': 'Connecticut', 'Del.': 'Delaware', 'Fla.': 'Florida', 'Ga.': 'Georgia', 
    'Hawaii': 'Hawaii', 'Idaho': 'Idaho', 'Ill.': 'Illinois', 'Ind.': 'Indiana', 'Iowa': 'Iowa', 
    'Kans.': 'Kansas', 'Ky.': 'Kentucky', 'La.': 'Louisiana', 'Maine': 'Maine', 'Md.': 'Maryland', 
    'Mass.': 'Massachusetts', 'Mich.': 'Michigan', 'Minn.': 'Minnesota', 'Miss.': 'Mississippi', 
    'Mo.': 'Missouri', 'Mont.': 'Montana', 'Nebr.': 'Nebraska', 'Nev.': 'Nevada', 'N.H.': 'New Hampshire', 
    'N.J.': 'New Jersey', 'N.M.': 'New Mexico', 'N.Y.': 'New York', 'N.C.': 'North Carolina', 
    'N.D.': 'North Dakota', 'Ohio': 'Ohio', 'Okla.': 'Oklahoma', 'Ore.': 'Oregon', 'Pa.': 'Pennsylvania', 
    'R.I.': 'Rhode Island', 'S.C.': 'South Carolina', 'S.D.': 'South Dakota', 'Tenn.': 'Tennessee', 
    'Tex.': 'Texas', 'Utah': 'Utah', 'Vt.': 'Vermont', 'Va.': 'Virginia', 'Wash.': 'Washington', 
    'W.Va.': 'West Virginia', 'Wis.': 'Wisconsin', 'Wyo.': 'Wyoming', 'D.C.': 'District of Columbia'
}

# States that tax only investment income (not wages) in specific years.
# New Hampshire I&D tax repealed from TY 2025; Tennessee Hall tax repealed from TY 2021.
INTEREST_DIVIDENDS_ONLY_YEARS = {
    'New Hampshire': list(range(2015, 2025)),
    'Tennessee': list(range(2015, 2021)),
}
# Washington capital gains excise tax effective from TY 2022.
CAPITAL_GAINS_ONLY_YEARS = {
    'Washington': list(range(2022, 2027)),
}

def clean_val(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip().replace('\xa0', '')
        if not v:
            return None
    return v

def parse_credit_or_amount(v):
    if v is None:
        return 0, False
    if isinstance(v, (int, float)):
        return v, False
    v_str = str(v).lower()
    if 'n.a' in v_str:
        return 0, False
    
    is_credit = 'credit' in v_str
    # Extract number (remove $ and commas)
    nums = re.findall(r'[\d,]+\.?\d*', v_str)
    if nums:
        num_str = nums[0].replace(',', '')
        try:
            return float(num_str) if '.' in num_str else int(num_str), is_credit
        except ValueError:
            return 0, is_credit
    return 0, is_credit

def extract_footnotes(s):
    s = str(s)
    match = re.search(r'\((.*?)\)', s)
    footnotes = []
    if match:
        footnotes_str = match.group(1)
        footnotes = [f.strip() for f in footnotes_str.split(',')]
    state_str = re.sub(r'\(.*?\)', '', s).strip()
    return state_str, footnotes

def parse_rate(v):
    if v is None: return None
    if isinstance(v, (int, float)):
        if v > 1: return float(v) / 100.0
        return float(v)
    v_str = str(v).strip().lower()
    if v_str == 'none': return None
    has_percent = '%' in v_str
    v_str = v_str.replace('%', '')
    try:
        # Check if there are footnote strings inside like '2.9% (b)'
        m = re.match(r'([\d.]+)', v_str)
        if m:
            val = float(m.group(1))
            if has_percent or val > 1:
                return val / 100.0
            return val
    except ValueError:
        pass
    return None

def parse_bracket(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    v_str = str(v).strip()
    m = re.search(r'([\d,]+\\.?\\d*)', v_str)
    if m:
        return float(m.group(1).replace(',', ''))
    return None

def parse_sheet(ws, year: int):
    """Parse one year sheet into (state_data, footnotes_dict)."""
    data = {}
    footnotes_dict = {}

    current_state_data = None
    row_idx = 3  # 1-based, first two are headers
    max_row = ws.max_row

    while row_idx <= max_row:
        row = [clean_val(ws.cell(row=row_idx, column=col).value) for col in range(1, 13)]
        col0 = str(row[0] or '').strip()
        
        # Check for footnotes section (code followed by descriptive text).
        # A bare "(code)" with no text is a state continuation footnote (e.g. Louisiana "(d)").
        if re.match(r'^\([a-z]+\)', col0) and not any(row[1:]):
            match = re.match(r'^\(([a-z]+)\)\s*(.*)', col0)
            if match and match.group(2).strip():
                fn_code = match.group(1)
                fn_text = match.group(2).strip()
                footnotes_dict[fn_code] = fn_text
                row_idx += 1
                continue

        # Check for empty row
        if not any(row):
            row_idx += 1
            continue
            
        if current_state_data and not any(row[0:12]) and not col0.startswith('('):
            row_idx += 1
            continue

        is_new_state = False
        
        if col0:
            state_part, fns = extract_footnotes(col0)
            if state_part in STATE_MAP:
                is_new_state = True
                current_state = STATE_MAP[state_part]
                current_state_data = {
                    "state_name": current_state,
                    "state_abbr": state_part,
                    "has_income_tax": True,
                    "is_capital_gains_only": current_state in CAPITAL_GAINS_ONLY_YEARS and year in CAPITAL_GAINS_ONLY_YEARS[current_state],
                    "is_interest_dividends_only": current_state in INTEREST_DIVIDENDS_ONLY_YEARS and year in INTEREST_DIVIDENDS_ONLY_YEARS[current_state],
                    "brackets_single": [],
                    "brackets_married": [],
                    "standard_deduction_single": 0,
                    "standard_deduction_married": 0,
                    "standard_deduction_is_credit": False,
                    "personal_exemption_single": 0,
                    "personal_exemption_married": 0,
                    "personal_exemption_dependent": 0,
                    "exemption_is_credit": False,
                    "dependent_exemption_is_credit": False,
                    "footnotes": fns,
                    "notes": ""
                }
                data[current_state] = current_state_data
            elif current_state_data and fns:
                # Continuation row carrying footnote codes (may also carry bracket data)
                current_state_data['footnotes'].extend(fns)
                if not any(row[1:12]):
                    row_idx += 1
                    continue

        if current_state_data:
            # Deductions and exemptions usually on the first row of state
            if is_new_state:
                sd_s, sd_s_credit = parse_credit_or_amount(row[7])
                sd_m, sd_m_credit = parse_credit_or_amount(row[8])
                current_state_data["standard_deduction_single"] = sd_s
                current_state_data["standard_deduction_married"] = sd_m
                current_state_data["standard_deduction_is_credit"] = bool(sd_s_credit or sd_m_credit)
                if sd_s_credit or sd_m_credit:
                    current_state_data["notes"] += "Standard deduction is a credit. "
                    
                ex_s, ex_s_credit = parse_credit_or_amount(row[9])
                ex_m, ex_m_credit = parse_credit_or_amount(row[10])
                ex_d, ex_d_credit = parse_credit_or_amount(row[11])
                
                current_state_data["personal_exemption_single"] = ex_s
                current_state_data["personal_exemption_married"] = ex_m
                current_state_data["personal_exemption_dependent"] = ex_d
                
                if ex_s_credit or ex_m_credit:
                    current_state_data["exemption_is_credit"] = True
                if ex_d_credit:
                    current_state_data["dependent_exemption_is_credit"] = True

            rate_s = row[1]
            bracket_s = row[3]
            rate_m = row[4]
            bracket_m = row[6]

            if rate_s is not None and str(rate_s).strip().lower() == 'none':
                current_state_data["has_income_tax"] = False
                
            if current_state_data["has_income_tax"]:
                # Single
                if rate_s is not None and str(rate_s).strip().lower() != 'none':
                    rs_val = parse_rate(rate_s)
                    bs_val = parse_bracket(bracket_s)
                    if rs_val is not None and bs_val is not None:
                        current_state_data["brackets_single"].append({"rate": rs_val, "min_income": bs_val, "max_income": None})
                
                # Married
                if rate_m is not None and str(rate_m).strip().lower() != 'none':
                    rm_val = parse_rate(rate_m)
                    bm_val = parse_bracket(bracket_m)
                    if rm_val is not None and bm_val is not None:
                        current_state_data["brackets_married"].append({"rate": rm_val, "min_income": bm_val, "max_income": None})

        row_idx += 1

    # Post-process: fill max_income for each bracket
    for state, s_data in data.items():
        b_s = s_data["brackets_single"]
        for i in range(len(b_s)):
            if i < len(b_s) - 1:
                b_s[i]["max_income"] = b_s[i+1]["min_income"]
                
        b_m = s_data["brackets_married"]
        for i in range(len(b_m)):
            if i < len(b_m) - 1:
                b_m[i]["max_income"] = b_m[i+1]["min_income"]

    return data, footnotes_dict


wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)

summary = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    year = int(sheet_name)
    data, footnotes = parse_sheet(ws, year)
    year_str = str(sheet_name)
    
    out_tax = DATA_DIR / f'tax_data_{year_str}.json'
    out_fn = DATA_DIR / f'footnotes_{year_str}.json'
    with open(out_tax, 'w') as f:
        json.dump(data, f, indent=2)
    with open(out_fn, 'w') as f:
        json.dump(footnotes, f, indent=2)
    
    summary[year_str] = {
        'states': len(data),
        'footnotes': len(footnotes),
        'no_tax': sorted([s for s, v in data.items() if not v['has_income_tax']]),
        'cg_only': [s for s, v in data.items() if v['is_capital_gains_only']],
    }

for year in sorted(summary.keys()):
    s = summary[year]
    print(f"{year}: states={s['states']}, footnotes={s['footnotes']}, no_tax={s['no_tax']}, cg_only={s['cg_only']}")
